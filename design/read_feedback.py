# design/デザイン案.xlsx の「修正コメント（ここに記入）」列を読み、シートごとに一覧で出す。
#   uv run --with openpyxl python design/read_feedback.py          … 一覧（Markdown）
#   uv run --with openpyxl python design/read_feedback.py --json   … JSON
# Claude はこの出力を見て spec.json / text.json を直し、make_assets.py → build_xlsx.py で作り直す。
import json
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, 'デザイン案.xlsx')
COMMENT_HEAD = '修正コメント（ここに記入）'


def read():
    wb = load_workbook(XLSX)
    out = []
    for ws in wb.worksheets:
        head = [c.value for c in ws[1]]
        if COMMENT_HEAD not in head or 'ID' not in head:
            continue
        ci, ii = head.index(COMMENT_HEAD), head.index('ID')
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[ii] and row[ci] and str(row[ci]).strip():
                ctx = [str(v) for k, v in enumerate(row) if v and k not in (ii, ci) and head[k] in ('案の名前', '見出し', '項目', '見出し')]
                out.append({'sheet': ws.title, 'id': str(row[ii]), 'target': ' / '.join(ctx), 'comment': str(row[ci]).strip()})
    return out


if __name__ == '__main__':
    items = read()
    if '--json' in sys.argv:
        print(json.dumps(items, ensure_ascii=False, indent=2))
    elif not items:
        print('コメントはまだありません。')
    else:
        for it in items:
            print(f"- [{it['sheet']}] {it['id']}（{it['target']}）: {it['comment']}")
